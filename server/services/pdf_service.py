import uuid
from pathlib import Path
import pymupdf as fitz
from pdf2docx import Converter
from pypdf import PdfReader, PdfWriter
from services.cleanup_service import TEMP_DIR

def convert_pdf_to_docx(pdf_bytes: bytes, target_font: str = "Times New Roman") -> Path:
    unique_id = str(uuid.uuid4())
    input_path = TEMP_DIR / f"{unique_id}.pdf"
    output_path = TEMP_DIR / f"{unique_id}.docx"

    with open(input_path, "wb") as f:
        f.write(pdf_bytes)

    cv = Converter(str(input_path))
    cv.convert(str(output_path), start=0, end=None)
    cv.close()

    # Normalize font in generated docx to Times New Roman or general font
    try:
        from docx import Document
        doc = Document(str(output_path))
        for p in doc.paragraphs:
            for run in p.runs:
                run.font.name = target_font
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = target_font
        doc.save(str(output_path))
    except Exception:
        pass

    input_path.unlink(missing_ok=True)
    return output_path

def convert_pdf_to_txt(pdf_bytes: bytes) -> Path:
    unique_id = str(uuid.uuid4())
    output_path = TEMP_DIR / f"{unique_id}.txt"

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text_parts = []
    for page in doc:
        text_parts.append(page.get_text("text"))
    doc.close()

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n--- Halaman Dokumen ---\n\n".join(text_parts))

    return output_path

def convert_pdf_to_image(pdf_bytes: bytes, fmt: str = "jpg") -> Path:
    unique_id = str(uuid.uuid4())
    output_path = TEMP_DIR / f"{unique_id}.{fmt}"

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    if len(doc) > 0:
        page = doc[0]
        pix = page.get_pixmap(dpi=150)
        pix.save(str(output_path))
    doc.close()
    return output_path

def compress_pdf(pdf_bytes: bytes, level: str = "recommended") -> Path:
    unique_id = str(uuid.uuid4())
    output_path = TEMP_DIR / f"compressed_{unique_id}.pdf"

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")

    deflate = True
    garbage = 4 if level in ["extreme", "recommended"] else 2
    clean = True if level == "extreme" else False

    doc.save(
        str(output_path),
        garbage=garbage,
        deflate=deflate,
        clean=clean,
    )
    doc.close()
    return output_path

def protect_pdf(pdf_bytes: bytes, password: str) -> Path:
    unique_id = str(uuid.uuid4())
    output_path = TEMP_DIR / f"protected_{unique_id}.pdf"

    import io
    reader = PdfReader(io.BytesIO(pdf_bytes))
    writer = PdfWriter()

    for page in reader.pages:
        writer.add_page(page)

    writer.encrypt(password)

    with open(output_path, "wb") as f:
        writer.write(f)

    return output_path

def unlock_pdf(pdf_bytes: bytes, password: str) -> Path:
    unique_id = str(uuid.uuid4())
    output_path = TEMP_DIR / f"unlocked_{unique_id}.pdf"

    import io
    reader = PdfReader(io.BytesIO(pdf_bytes))
    if reader.is_encrypted:
        reader.decrypt(password)

    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)

    with open(output_path, "wb") as f:
        writer.write(f)

    return output_path
